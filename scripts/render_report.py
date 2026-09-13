#!/usr/bin/env python3
"""Render a reviewed paper-report JSON to continuous HTML and/or simple PPTX."""

import argparse
import base64
import html
import json
import mimetypes
from pathlib import Path


def metadata_text(value, separator: str) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list) and all(isinstance(item, str) for item in value):
        return separator.join(item.strip() for item in value if item.strip())
    raise ValueError("paper metadata fields must be strings or lists of strings")


def ppt_text(value) -> str:
    return str(value).rstrip("。；;，,")


def load_report(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not isinstance(data.get("sections"), list) or not data["sections"]:
        raise ValueError("Report must contain a nonempty sections list")
    paper = data.get("paper")
    if not isinstance(paper, dict) or not paper.get("title"):
        raise ValueError("paper.title is required")
    for field in ("authors", "affiliations"):
        if field in paper:
            metadata_text(paper[field], " · ")
    if not data.get("report_title"):
        raise ValueError("report_title is required")
    if data.get("language", "zh") not in {"zh", "en"}:
        raise ValueError("language must be zh or en")
    for number, section in enumerate(data["sections"], 1):
        if not isinstance(section, dict):
            raise ValueError(f"Section {number} must be an object")
        kind = section.get("type")
        if kind not in {"paper", "analysis", "question"}:
            raise ValueError(f"Section {number} has invalid type")
        if not section.get("title"):
            raise ValueError(f"Section {number} needs a title")
        if kind == "paper" and section.get("takeaway") and not section.get("takeaway_source"):
            raise ValueError(f"Section {number} paper takeaway needs a source")
        points = section.get("points", [])
        if not isinstance(points, list):
            raise ValueError(f"Section {number} points must be a list")
        for index, point in enumerate(points, 1):
            if not isinstance(point, dict) or not point.get("text"):
                raise ValueError(f"Section {number} point {index} needs text")
            if kind == "paper" and not point.get("source"):
                raise ValueError(f"Section {number} point {index} needs a source")
        figure = section.get("figure")
        if figure is not None:
            if not isinstance(figure, dict) or not all(figure.get(key) for key in ("path", "caption", "source")):
                raise ValueError(f"Section {number} figure needs path, caption, source")
            image_path = Path(figure["path"])
            if not image_path.is_absolute():
                image_path = path.parent / image_path
            if not image_path.is_file():
                raise FileNotFoundError(f"Section {number} figure missing: {image_path}")
            figure["_resolved_path"] = image_path.resolve()
    return data


def validate_speaker_script(data: dict) -> dict:
    script = data.get("speaker_script")
    if not isinstance(script, dict):
        raise ValueError("speaker_script is required for PPTX output")
    target_minutes = script.get("target_minutes", 10)
    if not isinstance(target_minutes, int) or not 1 <= target_minutes <= 60:
        raise ValueError("speaker_script.target_minutes must be an integer from 1 to 60")
    if not isinstance(script.get("opening"), str) or not script["opening"].strip():
        raise ValueError("speaker_script.opening is required")
    entries = script.get("sections")
    if not isinstance(entries, list) or len(entries) != len(data["sections"]):
        raise ValueError("speaker_script.sections must contain one entry for every report section")
    total_seconds = 0
    for number, (section, entry) in enumerate(zip(data["sections"], entries), 1):
        if not isinstance(entry, dict) or entry.get("title") != section["title"]:
            raise ValueError(f"speaker_script section {number} must match the report section title")
        if not isinstance(entry.get("script"), str) or not entry["script"].strip():
            raise ValueError(f"speaker_script section {number} needs script text")
        seconds = entry.get("duration_seconds")
        if not isinstance(seconds, int) or not 15 <= seconds <= 180:
            raise ValueError(f"speaker_script section {number} needs duration_seconds from 15 to 180")
        total_seconds += seconds
    qa = script.get("qa")
    if not isinstance(qa, list) or len(qa) < 3:
        raise ValueError("speaker_script.qa needs at least three question-and-answer entries")
    for number, item in enumerate(qa, 1):
        if not isinstance(item, dict) or not all(isinstance(item.get(key), str) and item[key].strip() for key in ("question", "answer")):
            raise ValueError(f"speaker_script QA {number} needs question and answer")
    opening_seconds = script.get("opening_duration_seconds", 40)
    if not isinstance(opening_seconds, int) or not 15 <= opening_seconds <= 120:
        raise ValueError("speaker_script.opening_duration_seconds must be from 15 to 120")
    expected_seconds = target_minutes * 60
    if not expected_seconds - 100 <= total_seconds + opening_seconds <= expected_seconds + 100:
        raise ValueError("speaker_script timings must be close to target_minutes")
    return script


def speaker_script_markdown(data: dict) -> str:
    script = validate_speaker_script(data)
    lines = [
        f"# {data['report_title']}｜讲稿",
        "",
        f"> 预计时长：约 {script['target_minutes']} 分钟（不含最后问答）",
        "",
        f"## 开场（约 {script.get('opening_duration_seconds', 40)} 秒）",
        "",
        script["opening"].strip(),
    ]
    for number, entry in enumerate(script["sections"], 1):
        lines.extend([
            "",
            f"## 第 {number} 页：{entry['title']}（约 {entry['duration_seconds']} 秒）",
            "",
            entry["script"].strip(),
        ])
    lines.extend(["", "## 可能的提问与回答"])
    for number, item in enumerate(script["qa"], 1):
        lines.extend(["", f"### Q{number}. {item['question'].strip()}", "", item["answer"].strip()])
    return "\n".join(lines) + "\n"


def html_document(data: dict) -> str:
    zh = data.get("language", "zh") == "zh"
    esc = lambda value: html.escape(str(value), quote=True)
    styles = """
    :root{--bg:#f8f9fb;--paper:#fff;--ink:#172338;--muted:#5a687b;--line:#dce2ea;--accent:#254b78}
    *{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:var(--bg);color:var(--ink);font:17px/1.65 'Microsoft YaHei','Noto Sans CJK SC',Arial,sans-serif}
    .wrap{max-width:1080px;margin:auto;padding:56px 44px 96px}header{border-bottom:4px solid var(--accent);padding:18px 0 38px;margin-bottom:0}
    h1{font-size:2.2rem;line-height:1.25;margin:0 0 16px}h2{font-size:1.55rem;line-height:1.35;margin:0 0 18px}p{margin:0 0 16px}
    .meta{color:var(--muted);font-size:.84rem}.paper-title{font-size:1.06rem;color:var(--accent)}.authors{margin-top:16px;color:var(--ink);font-size:.96rem}.affiliations{white-space:pre-line;color:var(--muted);font-size:.84rem;margin-top:5px}
    section{border-bottom:1px solid var(--line);padding:52px 0 46px;scroll-margin-top:18px}
    .takeaway{font-size:1.18rem;font-weight:600;border-left:4px solid var(--accent);padding:8px 0 8px 18px;margin-bottom:22px}
    ul{padding-left:24px;margin:0 0 18px}li{margin:8px 0}figure{margin:28px 0 4px;padding:18px;background:var(--paper);border:1px solid var(--line)}figure img{max-width:100%;max-height:580px;display:block;margin:auto;object-fit:contain}figcaption{text-align:center;margin-top:10px;color:var(--muted);font-size:.86rem}
    @media(max-width:650px){.wrap{padding:28px 20px 64px}header{padding-bottom:28px}section{padding:38px 0}h1{font-size:1.75rem}}
    @media print{body{background:#fff}.wrap{max-width:none;padding:0}header,section{border:0;border-bottom:1px solid #bbb;break-inside:avoid;padding:18px 0}}
    """
    title = esc(data["report_title"])
    paper = data["paper"]
    authors = esc(metadata_text(paper.get("authors"), " · "))
    affiliations = esc(metadata_text(paper.get("affiliations"), "\n"))
    out = [f'<!doctype html><html lang="{"zh-CN" if zh else "en"}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title><style>{styles}</style></head><body><main class="wrap">']
    out.append(f'<header><h1>{title}</h1><p class="paper-title">{esc(paper["title"])}</p>')
    if authors:
        out.append(f'<p class="authors">{authors}</p>')
    if affiliations:
        out.append(f'<p class="affiliations">{affiliations}</p>')
    out.append(f'<p class="meta">{esc(paper.get("venue", ""))}</p></header>')
    for index, section in enumerate(data["sections"], 1):
        out.append(f'<section id="section-{index}"><h2>{esc(section["title"])}</h2>')
        if section.get("takeaway"):
            out.append(f'<p class="takeaway">{esc(section["takeaway"])}</p>')
        if section.get("points"):
            out.append('<ul>')
            for point in section["points"]:
                out.append(f'<li>{esc(point["text"])}</li>')
            out.append('</ul>')
        if section.get("figure"):
            figure = section["figure"]
            image_path = figure["_resolved_path"]
            mime = mimetypes.guess_type(image_path.name)[0] or "image/png"
            encoded = base64.b64encode(image_path.read_bytes()).decode("ascii")
            out.append(f'<figure><img src="data:{mime};base64,{encoded}" alt="{esc(figure["caption"])}"><figcaption>{esc(figure["caption"])}</figcaption></figure>')
        out.append('</section>')
    out.append('</main></body></html>')
    return ''.join(out)


def render_pptx(data: dict, target: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    for number, section in enumerate(data["sections"], 1):
        if len(section.get("points", [])) > 5 or any(len(point["text"]) > 120 for point in section.get("points", [])):
            raise ValueError(f"Section {number} is too dense for PPTX; split or shorten its points")
        if len(section.get("takeaway", "")) > 120:
            raise ValueError(f"Section {number} takeaway is too long for PPTX")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    navy, ink, muted = RGBColor(37, 75, 120), RGBColor(23, 35, 56), RGBColor(88, 103, 123)
    zh = data.get("language", "zh") == "zh"
    font = "Microsoft YaHei" if zh else "Aptos"

    def text_box(slide, x, y, w, h, value, size=20, color=ink, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.02)
        frame.margin_top = frame.margin_bottom = Inches(.01)
        for row, line in enumerate(str(value).split("\n")):
            paragraph = frame.paragraphs[0] if row == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = font
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = color
            paragraph.space_after = Pt(3)
        return shape

    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    authors = metadata_text(data["paper"].get("authors"), " · ")
    affiliations = metadata_text(data["paper"].get("affiliations"), "\n")
    text_box(title_slide, .75, .72, 11.8, .9, data["report_title"], 30, navy, True)
    text_box(title_slide, .78, 1.82, 11.6, 1.08, data["paper"]["title"], 23, ink, True)
    if authors:
        text_box(title_slide, .8, 3.18, 11.8, .78, authors, 14, ink)
    if affiliations:
        text_box(title_slide, .8, 4.05, 11.8, 1.62, affiliations, 10, muted)
    subtitle = data["paper"].get("venue", "")
    text_box(title_slide, .8, 6.32, 11.8, .36, subtitle, 12, muted)
    for section in data["sections"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text_box(slide, .62, .28, 11.7, .65, section["title"], 29, navy, True)
        has_figure = bool(section.get("figure"))
        content_w = 5.65 if has_figure else 11.7
        y = 1.25
        if section.get("takeaway"):
            text_box(slide, .68, y, content_w, 1.24, ppt_text(section["takeaway"]), 20, ink, True)
            y += 1.38
        points = section.get("points", [])
        remaining = max(1.0, 6.78 - y)
        per_point = min(1.18, remaining / max(len(points), 1))
        point_size = 17 if len(points) <= 4 else 15
        for point in points:
            marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(.78), Inches(y + .16), Inches(.10), Inches(.10))
            marker.fill.solid()
            marker.fill.fore_color.rgb = navy
            marker.line.fill.background()
            text_box(slide, .98, y, content_w - .3, per_point, ppt_text(point["text"]), point_size)
            y += per_point
        if has_figure:
            figure = section["figure"]
            from PIL import Image
            with Image.open(figure["_resolved_path"]) as picture:
                width, height = picture.size
            max_w, max_h = 5.6, 4.75
            ratio = min(max_w / width, max_h / height)
            draw_w, draw_h = width * ratio, height * ratio
            slide.shapes.add_picture(str(figure["_resolved_path"]), Inches(7.15 + (max_w - draw_w) / 2), Inches(1.5 + (max_h - draw_h) / 2), width=Inches(draw_w), height=Inches(draw_h))
            text_box(slide, 7.18, 6.42, 5.55, .5, figure["caption"], 10, muted)
        if section.get("speaker_notes"):
            frame = slide.notes_slide.notes_text_frame
            frame.text = section["speaker_notes"]
    prs.save(target)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("report", type=Path)
    parser.add_argument("--format", choices=["html", "pptx", "both"], default="html")
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    report = load_report(args.report.resolve())
    args.out.mkdir(parents=True, exist_ok=True)
    if args.format in {"pptx", "both"}:
        target = args.out / "report.pptx"
        render_pptx(report, target)
        print(target)
        script_target = args.out / "speaker-script.md"
        script_target.write_text(speaker_script_markdown(report), encoding="utf-8")
        print(script_target)
    if args.format in {"html", "both"}:
        target = args.out / "report.html"
        target.write_text(html_document(report), encoding="utf-8")
        print(target)


if __name__ == "__main__":
    main()
