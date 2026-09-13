import json
import subprocess
import tempfile
import unittest
from pathlib import Path
import sys

from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from inspect_pdf import extract
from render_report import html_document, load_report, render_pptx, speaker_script_markdown


class WorkflowTests(unittest.TestCase):
    def test_pdf_inspection_and_both_formats(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pdf = root / "paper.pdf"
            embedded = Image.new("RGB", (420, 260), "navy")
            c = canvas.Canvas(str(pdf))
            c.drawString(72, 740, "Figure 1: Method overview")
            c.drawString(72, 700, "The method improves accuracy on the test set.")
            c.drawImage(ImageReader(embedded), 72, 350, width=300, height=186)
            c.save()
            inspected = extract(pdf, root / "work")
            self.assertEqual(inspected["pages"], 1)
            self.assertEqual(inspected["figure_candidates"], 1)
            pages = json.loads(Path(inspected["text"]).read_text(encoding="utf-8"))
            self.assertIn("Method overview", pages["pages"][0]["text"])
            figures = json.loads(Path(inspected["figure_index"]).read_text(encoding="utf-8"))
            self.assertTrue((root / "work" / figures["figures"][0]["path"]).is_file())
            cropped = extract(pdf, root / "cropped", crop=[1, 65, 230, 385, 500])
            self.assertEqual(cropped["figure_candidates"], 2)

            image = root / "figure.png"
            Image.new("RGB", (640, 360), "white").save(image)
            data = {
                "report_title": "论文组会汇报",
                "paper": {"title": "A Method Paper", "authors": ["A. Author¹", "B. Author²"], "affiliations": ["¹ University A", "² Institute B"], "venue": "Test 2026", "pdf": str(pdf)},
                "language": "zh", "duration_minutes": 20,
                "speaker_script": {
                    "target_minutes": 6,
                    "opening_duration_seconds": 30,
                    "opening": "今天我们用十分钟看看这篇论文解决什么问题，以及它怎样让检索更可靠。",
                    "sections": [
                        {"title": "方法", "duration_seconds": 180, "script": "这一页先讲方法的直觉。它把问题拆成检索和回答两个简单步骤。"},
                        {"title": "讨论", "duration_seconds": 180, "script": "这一页提醒我们，结果还需要在更多场景里验证。"}
                    ],
                    "qa": [
                        {"question": "它解决了什么问题？", "answer": "它帮助系统找到更贴近问题的信息。"},
                        {"question": "为什么要检索？", "answer": "检索可以给模型补充外部证据。"},
                        {"question": "还有什么限制？", "answer": "还需要更多真实场景的测试。"}
                    ]
                },
                "sections": [
                    {"type": "paper", "title": "方法", "takeaway": "方法有效", "takeaway_source": "p. 1", "points": [{"text": "准确率提升", "source": "p. 1"}], "figure": {"path": "figure.png", "caption": "Fig. 1", "source": "p. 1"}, "speaker_notes": "解释方法。"},
                    {"type": "analysis", "title": "讨论", "takeaway": "需要更广泛的验证", "points": [{"text": "可讨论其他数据集。"}]}
                ]
            }
            report_path = root / "report.json"
            report_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            report = load_report(report_path)
            html_text = html_document(report)
            self.assertEqual(html_text.count("<section id="), 2)
            self.assertIn('id="section-2"', html_text)
            self.assertNotIn("is-active", html_text)
            self.assertNotIn("<nav", html_text)
            self.assertNotIn("<ol>", html_text)
            self.assertNotIn("讲解备注", html_text)
            self.assertNotIn("Speaker notes", html_text)
            self.assertNotIn("<details", html_text)
            self.assertIn("A. Author¹ · B. Author²", html_text)
            self.assertIn("¹ University A\n² Institute B", html_text)
            self.assertIn("data:image/png;base64,", html_text)
            self.assertNotIn("论文内容", html_text)
            self.assertNotIn("汇报者分析", html_text)
            self.assertNotIn("讨论问题", html_text)
            self.assertNotIn("p. 1", html_text)
            self.assertNotIn("20 min", html_text)
            markdown = speaker_script_markdown(report)
            self.assertIn("预计时长：约 6 分钟", markdown)
            self.assertIn("第 1 页：方法", markdown)
            self.assertIn("可能的提问与回答", markdown)
            target = root / "report.pptx"
            render_pptx(report, target)
            from pptx import Presentation
            prs = Presentation(target)
            self.assertEqual(len(prs.slides), 3)
            self.assertIn("解释方法", prs.slides[1].notes_slide.notes_text_frame.text)
            title_text = "\n".join(shape.text for shape in prs.slides[0].shapes if hasattr(shape, "text"))
            self.assertIn("A. Author¹ · B. Author²", title_text)
            self.assertIn("¹ University A", title_text)
            self.assertNotIn("20 min", title_text)
            slide_text = "\n".join(shape.text for shape in prs.slides[1].shapes if hasattr(shape, "text"))
            self.assertNotIn("来源：", slide_text)
            self.assertNotIn("论文内容", slide_text)
            self.assertNotIn("汇报者分析", slide_text)
            self.assertNotIn("讨论问题", slide_text)
            self.assertNotIn("p. 1", slide_text)
            subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "render_report.py"), str(report_path), "--format", "both", "--out", str(root / "output")], check=True, capture_output=True, text=True)
            self.assertTrue((root / "output" / "report.html").is_file())
            self.assertTrue((root / "output" / "report.pptx").is_file())
            self.assertTrue((root / "output" / "speaker-script.md").is_file())

    def test_paper_claim_without_source_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.json"
            path.write_text(json.dumps({"report_title": "T", "paper": {"title": "P"}, "sections": [{"type": "paper", "title": "Claim", "points": [{"text": "Unsupported"}]}]}), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "needs a source"):
                load_report(path)

    def test_pptx_rejects_unreadable_density(self):
        with tempfile.TemporaryDirectory() as tmp:
            data = {"report_title": "T", "paper": {"title": "P"}, "language": "en", "sections": [{"type": "analysis", "title": "Too much", "points": [{"text": "x"} for _ in range(6)]}]}
            with self.assertRaisesRegex(ValueError, "too dense"):
                render_pptx(data, Path(tmp) / "report.pptx")


if __name__ == "__main__":
    unittest.main()
